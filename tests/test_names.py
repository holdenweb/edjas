"""Tests for listing what a workbook offers as references (``edjas.names``).

The interesting cases are the unusable ones. Real statistical workbooks are last
year's file with the figures changed, so they accumulate names whose cells were
deleted years ago and names that hold constants rather than ranges. Reporting those
as what they are -- rather than as a damaged sheet, or by omitting them -- is most of
the value of this listing.
"""

import io

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from edjas import main
from edjas.names import describe, report


def workbook_with(names=(), tables=(), sheet="Data"):
    book = Workbook()
    page = book.active
    page.title = sheet
    for row in range(1, 5):
        for column in range(1, 4):
            page.cell(row=row, column=column, value=f"r{row}c{column}")
    for name, target in names:
        book.defined_names.add(DefinedName(name, attr_text=target))
    for name, refs in tables:
        page.add_table(Table(displayName=name, ref=refs))
    return book


def _problems(found):
    return {e["name"]: e["problem"] for e in found["names"]}


# --- what can be used --------------------------------------------------------

def test_a_usable_name_reports_its_sheet_range_and_shape():
    found = describe(workbook_with(names=[("Block", "Data!$A$1:$C$4")]))
    entry, = found["names"]
    assert entry["sheet"] == "Data"
    assert entry["refs"] == "$A$1:$C$4"
    assert entry["shape"] == (4, 3)
    assert entry["problem"] is None


def test_tables_are_listed_with_the_sheet_they_belong_to():
    found = describe(workbook_with(tables=[("Sales_Table", "A1:C4")]))
    entry, = found["tables"]
    assert entry["name"] == "Sales_Table"
    assert entry["sheet"] == "Data"
    assert entry["shape"] == (4, 3)


def test_names_and_tables_are_reported_separately():
    """They are stored in different places, which is why both need listing."""
    found = describe(workbook_with(
        names=[("Block", "Data!$A$1:$B$2")], tables=[("T", "A1:C4")],
    ))
    assert [e["name"] for e in found["names"]] == ["Block"]
    assert [e["name"] for e in found["tables"]] == ["T"]


# --- what cannot, and why ----------------------------------------------------

def test_a_name_whose_cells_were_deleted_is_reported_as_such():
    """Excel keeps the name and rewrites its target to #REF!; the sheets are fine."""
    found = describe(workbook_with(names=[("Gone", "#REF!")]))
    assert "deleted" in _problems(found)["Gone"]


def test_a_name_holding_a_constant_is_not_mistaken_for_a_sheet_reference():
    """An array constant may contain a '!' inside a quoted string.

    Splitting on that '!' would report a missing sheet named '{"Excel Help', which is
    a confusing diagnosis of a real but quite different situation.
    """
    target = '{"Excel Help!1802";5;10;13}'
    found = describe(workbook_with(names=[("Leftover", target)]))
    assert _problems(found)["Leftover"] == "holds a constant, not a range"


def test_a_union_range_is_reported_as_multi_area():
    found = describe(workbook_with(names=[("U", "Data!$A$1,Data!$C$3")]))
    assert "multi-area" in _problems(found)["U"]


def test_a_name_pointing_at_a_missing_sheet_says_so():
    found = describe(workbook_with(names=[("Away", "Elsewhere!$A$1")]))
    assert "Elsewhere" in _problems(found)["Away"]


# --- the report --------------------------------------------------------------

def test_dead_names_are_counted_rather_than_listed_by_default():
    """A workbook copied forward for years can carry hundreds; listing them all
    buries whatever is usable."""
    book = workbook_with(names=[(f"Dead{n}", "#REF!") for n in range(30)])
    out = io.StringIO()
    report(book, "book.xlsx", verbose=False, out=out)
    text = out.getvalue()
    assert "30 unusable" in text and "--verbose" in text
    assert "Dead7" not in text


def test_verbose_lists_every_dead_name():
    book = workbook_with(names=[(f"Dead{n}", "#REF!") for n in range(30)])
    out = io.StringIO()
    report(book, "book.xlsx", verbose=True, out=out)
    text = out.getvalue()
    assert "Dead7" in text and "unusable:" not in text


def test_the_suggested_form_matches_the_shape():
    book = workbook_with(names=[
        ("One", "Data!$A$1"), ("Pair", "Data!$A$1:$B$4"), ("Grid", "Data!$A$1:$C$4"),
    ])
    out = io.StringIO()
    report(book, "book.xlsx", out=out)
    text = out.getvalue()
    assert "e.g.  One" in text            # a single cell is a bare reference
    assert "e.g.  {Pair}" in text          # two columns make an object
    assert "e.g.  [Grid | records]" in text


# --- the command line --------------------------------------------------------

def test_list_names_prints_a_listing(tmp_path, capsys):
    path = tmp_path / "book.xlsx"
    workbook_with(names=[("Block", "Data!$A$1:$C$4")]).save(path)
    main([str(path), "--list-names"])
    out = capsys.readouterr().out
    assert "Named ranges" in out and "Block" in out


def test_list_names_refuses_a_spec(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook_with().save(path)
    with pytest.raises(SystemExit):
        main([str(path), "spec.toml", "--list-names"])


def test_a_spec_is_still_required_without_the_flag(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook_with().save(path)
    with pytest.raises(SystemExit):
        main([str(path)])
